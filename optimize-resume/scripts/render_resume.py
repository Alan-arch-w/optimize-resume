#!/usr/bin/env python3
"""Render factual resume JSON in the bundled reference layout or a user DOCX.
Copy formatting into a fresh package; never copy candidate content or embeds.
"""
import argparse
import json
import math
import unicodedata
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor, Pt

DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / 'assets' / 'reference-layout.docx'


def set_portable_fonts(doc, cjk_font, latin_font):
    """Use explicit, consistent font declarations instead of theme-dependent fallbacks."""
    for root in (doc.element, doc.styles.element):
        for fonts in root.iter(qn('w:rFonts')):
            fonts.attrib.clear()
            for key, value in {'ascii': latin_font, 'hAnsi': latin_font,
                               'cs': latin_font, 'eastAsia': cjk_font,
                               'hint': 'eastAsia'}.items():
                fonts.set(qn('w:' + key), value)
    # Register the same fonts in fontTable; it is a separate OPC part.
    table = next(
        rel.target_part for rel in doc.part.rels.values()
        if rel.reltype.endswith('/fontTable')
    )
    from lxml import etree
    root = etree.fromstring(table.blob)
    for name, charset in ((cjk_font, '86'), (latin_font, '00')):
        if not any(node.get(qn('w:name')) == name for node in root):
            node = OxmlElement('w:font')
            node.set(qn('w:name'), name)
            for tag, value in (('charset', charset), ('family', 'swiss'), ('pitch', 'variable')):
                child = OxmlElement('w:' + tag)
                child.set(qn('w:val'), value)
                node.append(child)
            root.append(node)
    table._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def text_units(text):
    """Conservative width estimate; rendered PDF remains the acceptance gate."""
    return sum(1 if unicodedata.east_asian_width(c) in ('W', 'F') else
               .3 if c.isspace() else .6 if c.isupper() else .52 for c in text)


def layout_metrics(scale):
    return {
        'body_size': 9 + 1.5 * scale, 'body_line': 12 + 3 * scale,
        'body_after': 2 + 1.5 * scale,
        'name_size': 17 + 3 * scale, 'name_line': 21 + 3 * scale,
        'section_size': 11 + scale, 'section_line': 15 + 2 * scale,
        'section_before': 8 + 4 * scale, 'section_after': 4 + 2 * scale,
        'entry_size': 9.5 + 1.5 * scale, 'entry_line': 13 + 3 * scale,
        'entry_before': 4 + 3 * scale, 'entry_after': 2 + scale,
        'role_after': 2 + 2 * scale,
    }


def line_text(value):
    if isinstance(value, dict):
        label = value['label'].rstrip('：:')
        return (label + '：' if label else '') + value['text']
    return value


def estimate_height(data, width_pt, scale):
    m = layout_metrics(scale)
    blocks = []

    def block(kind, text='', first=False):
        if kind == 'name':
            size, height, before, after = m['name_size'], m['name_line'], 0, 3
        elif kind == 'section':
            size, height, before, after = m['section_size'], m['section_line'], m['section_before'], m['section_after']
        elif kind == 'entry':
            size, height, before, after = m['entry_size'], m['entry_line'], 0 if first else m['entry_before'], m['entry_after']
        else:
            size, height, before, after = m['body_size'], m['body_line'], 0, m['role_after'] if kind == 'role' else m['body_after']
        rows = max(1, math.ceil(text_units(text) * size / width_pt))
        blocks.append((rows * height, before, after))

    block('name', data['name'])
    if data.get('contact'):
        block('body', ' | '.join(data['contact']))
    if data.get('target'):
        block('body', '求职意向：' + data['target'])
    groups = ([{'title': '个人总结', 'items': [{'bullets': data['summary']}]}]
              if data.get('summary') else []) + data['sections']
    for group in groups:
        if not group['items']:
            continue
        block('section', group['title'])
        for index, item in enumerate(group['items']):
            if item.get('heading') or item.get('date'):
                block('entry', item.get('heading', '') + '  ' + item.get('date', ''), first=index == 0)
            for field in ('role', 'detail'):
                if item.get(field):
                    block('role', item[field])
            for value in item.get('bullets', []):
                block('body', line_text(value))
    total, previous_after = 0, 0
    for height, before, after in blocks:
        total += max(previous_after, before) + height
        previous_after = after
    return total + previous_after


def select_scale(data, section, density):
    if density != 'auto':
        return {'compact': 0, 'balanced': .5, 'airy': 1}[density]
    width = (section.page_width - section.left_margin - section.right_margin) / 12700
    height = (section.page_height - section.top_margin - section.bottom_margin) / 12700
    minimum = estimate_height(data, width, 0)
    pages = max(1, math.ceil(minimum / (height * .97)))
    # Leave room for font/word wrapping and keep-with-next pagination effects.
    target = height * pages * (.94 if pages == 1 else .88)
    for step in range(20, -1, -1):
        scale = step / 20
        if estimate_height(data, width, scale) <= target:
            return scale
    return 0


def validate_line(value):
    if isinstance(value, str):
        return
    if isinstance(value, dict) and all(isinstance(value.get(k), str) for k in ('label', 'text')):
        return
    raise ValueError('line must be a string or {label: string, text: string}')


def validate(data):
    if not isinstance(data, dict) or not isinstance(data.get('name'), str) or not data['name'].strip():
        raise ValueError('name must be a nonempty string')
    for field in ('contact', 'summary'):
        if not isinstance(data.get(field, []), list):
            raise ValueError(f'{field} must be a list')
    if any(not isinstance(x, str) for x in data.get('contact', [])):
        raise ValueError('contact must contain strings')
    if not isinstance(data.get('target', ''), str):
        raise ValueError('target must be a string')
    if not isinstance(data.get('sections'), list) or not data['sections']:
        raise ValueError('sections must be a nonempty list')
    for group in data['sections']:
        if not isinstance(group, dict) or not isinstance(group.get('title'), str) or not isinstance(group.get('items'), list):
            raise ValueError('section requires title and items')
        for item in group['items']:
            if not isinstance(item, dict):
                raise ValueError('item must be an object')
            for key in ('heading', 'date', 'detail', 'role'):
                if not isinstance(item.get(key, ''), str):
                    raise ValueError(f'{key} must be a string')
            if not isinstance(item.get('bullets', []), list):
                raise ValueError('bullets must be a list')
            for value in item.get('bullets', []):
                validate_line(value)
    for value in data.get('summary', []):
        validate_line(value)


def render(data, output, template=None, font=None, density='auto', scale=None, adapt_template=False, latin_font=None):
    validate(data)
    source = Document(template or DEFAULT_TEMPLATE)
    # The deliverable must name fonts available on the recipient's machine,
    # not a renderer-only font installed on this server.
    if font is None and template is None:
        font = '微软雅黑'
    if latin_font is None:
        latin_font = 'Arial' if font in ('微软雅黑', 'Microsoft YaHei', '宋体', 'SimSun') else font
    doc = Document()
    # Copy styles, not the source package, relationships, metadata or body.
    styles = doc.styles.element
    for child in list(styles):
        styles.remove(child)
    for child in source.styles.element:
        styles.append(deepcopy(child))
    s, ss = doc.sections[0], source.sections[0]
    for attr in ('page_width', 'page_height', 'top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(s, attr, getattr(ss, attr))
    width = s.page_width - s.left_margin - s.right_margin
    responsive = template is None or adapt_template
    scale = select_scale(data, s, density) if scale is None else scale
    metrics = layout_metrics(scale)

    def prototype(predicate):
        return next((p for p in source.paragraphs if predicate(p)), None)
    title = prototype(lambda p: p.style.name == 'Title')
    contact = prototype(lambda p: p.alignment == WD_ALIGN_PARAGRAPH.CENTER and p.style.name != 'Title')
    section = prototype(lambda p: p.style.name.lower() == 'heading 1')
    entry = prototype(lambda p: '\t' in p.text)
    body = prototype(lambda p: p.style.name == 'Body Text' and p.alignment != WD_ALIGN_PARAGRAPH.CENTER)
    run_formats = {}

    def add_text(p, text, bold=None):
        r = p.add_run(text)
        if p._p in run_formats:
            r._r.insert(0, deepcopy(run_formats[p._p]))
        if bold is not None:
            r.bold = bold
        if font:
            r.font.name = font
            r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font)
        r.font.color.rgb = RGBColor(0, 0, 0)
        return r

    def paragraph(proto, text='', bold=None, keep=False, style=None, kind='body', first=False):
        p = doc.add_paragraph(style=style)
        if proto is not None and proto._p.pPr is not None:
            if p._p.pPr is not None:
                p._p.remove(p._p.pPr)
            p._p.insert(0, deepcopy(proto._p.pPr))
            for node in list(p._p.pPr):
                if node.tag in {qn('w:numPr'), qn('w:sectPr')}:
                    p._p.pPr.remove(node)
        if proto is not None:
            sample = next((r for r in proto.runs if r.text.strip()), None)
            if sample is not None and sample._r.rPr is not None:
                props = deepcopy(sample._r.rPr)
                for node in list(props):
                    if node.tag not in {qn('w:rFonts'), qn('w:sz'), qn('w:szCs'), qn('w:spacing')}:
                        props.remove(node)
                run_formats[p._p] = props
        if responsive:
            pf = p.paragraph_format
            m = metrics
            size, leading, before, after = m['body_size'], m['body_line'], 0, m['body_after']
            if kind == 'name':
                size, leading, before, after = m['name_size'], m['name_line'], 0, 3
            elif kind == 'section':
                size, leading, before, after = m['section_size'], m['section_line'], m['section_before'], m['section_after']
            elif kind == 'entry':
                size, leading, before, after = m['entry_size'], m['entry_line'], 0 if first else m['entry_before'], m['entry_after']
            elif kind == 'role':
                after = m['role_after']
            pf.space_before, pf.space_after = Pt(before), Pt(after)
            pf.line_spacing_rule, pf.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(leading)
            # Override template run sizes while retaining its typeface and black text.
            props = run_formats.get(p._p, OxmlElement('w:rPr'))
            for tag in ('w:sz', 'w:szCs'):
                for node in list(props.findall(qn(tag))):
                    props.remove(node)
                node = OxmlElement(tag)
                node.set(qn('w:val'), str(round(size * 2)))
                props.append(node)
            run_formats[p._p] = props
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = keep
        p.paragraph_format.widow_control = True
        if text:
            add_text(p, text, bold)
        return p

    paragraph(title, data['name'], True, True, 'Title', kind='name')
    if data.get('contact'):
        paragraph(contact, ' | '.join(data['contact']), keep=True, style='Body Text')
    if data.get('target'):
        paragraph(contact, '求职意向：' + data['target'], keep=True, style='Body Text')

    def heading(text):
        p = paragraph(section, text, True, True, 'Heading 1', kind='section')
        # Convert the reference's floating separator to a stable paragraph rule.
        # No source drawings/relationships are copied into the output.
        if section is not None and section._p.xpath('.//w:drawing | .//w:pict') and not p._p.xpath('./w:pPr/w:pBdr'):
            border = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            for key, val in {'val': 'single', 'sz': '6', 'color': '4F81BD', 'space': '1'}.items():
                bottom.set(qn('w:' + key), val)
            border.append(bottom)
            p._p.get_or_add_pPr().append(border)
        return p

    def line(value):
        p = paragraph(body, style='Body Text')
        if isinstance(value, dict):
            if value['label']:
                add_text(p, value['label'].rstrip('：:') + '：', True)
            add_text(p, value['text'])
        else:
            label, sep, text = value.partition('：')
            if sep and 0 < len(label) <= 18:
                add_text(p, label + sep, True)
                add_text(p, text)
            else:
                add_text(p, value)

    if data.get('summary'):
        heading('个人总结')
        for value in data['summary']:
            line(value)
    for group in data['sections']:
        if not group['items']:
            continue
        heading(group['title'])
        for index, item in enumerate(group['items']):
            if item.get('heading') or item.get('date'):
                p = paragraph(entry, keep=bool(item.get('role') or item.get('detail') or item.get('bullets')), kind='entry', first=index == 0)
                p.paragraph_format.tab_stops.clear_all()
                p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
                add_text(p, item.get('heading', ''), True)
                if item.get('date'):
                    add_text(p, '\t' + item['date'], False)
            if item.get('role'):
                paragraph(body, item['role'], keep=bool(item.get('detail') or item.get('bullets')), style='Body Text', kind='role')
            if item.get('detail'):
                paragraph(body, item['detail'], keep=bool(item.get('bullets')), style='Body Text', kind='role')
            for value in item.get('bullets', []):
                line(value)
    if font:
        for style in doc.styles:
            if style.type == 1:
                style.font.name = font
                style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font)
        set_portable_fonts(doc, font, latin_font)
    for key in ('author', 'last_modified_by', 'subject', 'comments', 'keywords'):
        setattr(doc.core_properties, key, '')
    doc.core_properties.title = '简历'
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return {'density': density, 'scale': scale, 'responsive': responsive,
            'estimated_height_pt': round(estimate_height(data, width / 12700, scale), 1)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input_json', type=Path)
    parser.add_argument('output_docx', type=Path)
    parser.add_argument('--template', type=Path, help='User reference DOCX; inspect visually before use')
    parser.add_argument('--font', help='Recipient-compatible Chinese font; bundled layout defaults to 微软雅黑')
    parser.add_argument('--latin-font', help='Recipient-compatible Latin font; defaults to Arial for common Chinese fonts')
    parser.add_argument('--density', choices=('auto', 'compact', 'balanced', 'airy'), default='auto')
    parser.add_argument('--scale', type=float, help='Override density on a 0 (compact) to 1 (airy) scale after PDF inspection')
    parser.add_argument('--adapt-template', action='store_true', help='Adjust an external template rhythm only when authorized')
    args = parser.parse_args()
    with args.input_json.open(encoding='utf-8') as f:
        data = json.load(f)
    if args.scale is not None and not 0 <= args.scale <= 1:
        parser.error('--scale must be between 0 and 1')
    report = render(data, args.output_docx, args.template, args.font, args.density, args.scale, args.adapt_template, args.latin_font)
    print(json.dumps({'output': str(args.output_docx.resolve()), **report}, ensure_ascii=False))


if __name__ == '__main__':
    main()
